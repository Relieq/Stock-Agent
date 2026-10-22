"""Xuất một BCTC đã đọc ra file Excel (.xlsx) cho người dùng.

Mỗi báo cáo (BCĐKT, KQKD, LCTT) một trang tính: số để dạng số theo đúng đơn vị in trên BCTC (không đổi về đồng), kèm
cột bằng chứng của từng ô (cùng phân loại và màu với trang tài liệu trên web) và số trang trong file gốc. Trang "Nguồn"
ghi BCTC gốc, giờ công bố, trạng thái kiểm chứng và các ô người duyệt đã sửa tay, để mọi số trong file vẫn truy được về
nguồn ("không nguồn, không số"). Chỉ xuất BCTC đã kiểm chứng hoặc đã duyệt.
"""

from __future__ import annotations

import io
from datetime import timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services.api.data import _check, detail
from services.extraction.evidence import STATUS_LABELS
from services.extraction.numbers import parse_amount

VN = timezone(timedelta(hours=7))
SHEET_NAMES = {"BS": "Cân đối kế toán", "IS": "Kết quả kinh doanh", "CF": "Lưu chuyển tiền tệ"}
UNIT_NAMES = {1: "đồng", 1_000: "nghìn đồng", 1_000_000: "triệu đồng", 1_000_000_000: "tỷ đồng"}
# Cùng màu với trang tài liệu trên web (services/api/static/style.css)
FILLS = {"both": "CDEED6", "period": "E3F4E8", "sum": "EEF7F0", "none": "F1F2F4", "restated": "FFF1D6",
         "differs": "DFE9FF", "second": "EBE7F5", "second_differs": "FBE0C4", "image_ok": "DCEFE2", "conflict": "F7B9B9"}
SHORT = {"both": "khớp cộng tổng và giữa các kỳ", "period": "khớp giữa các kỳ", "sum": "khớp cộng tổng",
         "none": "chưa kiểm được", "restated": "khác kỳ trước do trình bày lại",
         "differs": "hai BCTC in khác nhau", "second": "chưa có phép kiểm, hai lần đọc trùng",
         "second_differs": "hai lần đọc ra hai số, cần xem ảnh", "image_ok": "đã đối chiếu ảnh",
         "conflict": "cần xem ảnh"}
BOLD = Font(bold=True)


def _when(t) -> str:
    return t.astimezone(VN).strftime("%H:%M %d/%m/%Y") if t else ""


def workbook(record: dict, evidence, info: dict) -> bytes:
    """`info`: mã, tên công ty, tiêu đề, kỳ, phạm vi, nguồn, link nguồn, giờ công bố sớm nhất, giờ xử lý xong, trạng
    thái, số ràng buộc đạt (xem app.excel_export)."""
    d = detail(record, evidence)
    sep = d["sep"]
    statements = _check(record)[1]
    wb = Workbook()
    wb.remove(wb.active)
    for t in d["tables"]:
        ws = wb.create_sheet(SHEET_NAMES.get(t["statement"], t["statement"]))
        cols = t["columns"]
        # Đơn vị hệ thống đã dùng để hiểu số (báo cáo không in dòng đơn vị thì lấy theo báo cáo khác cùng file)
        mult = int(statements[t["statement"]].multiplier) if t["statement"] in statements else None
        unit = UNIT_NAMES.get(mult, t["unit"] or "không rõ")
        ws.append([f"{info['ticker']} · {t['name']} · {info['period_text']} · đơn vị: {unit}"])
        ws["A1"].font = BOLD
        ws.append(["STT", "Mã", "Chỉ tiêu", "Thuyết minh", *[n for _, n in cols], *[f"{n}: bằng chứng" for _, n in cols],
                   "Trang"])
        for c in ws[2]:
            c.font = BOLD
        for r in t["rows"]:
            nums = []
            for k, _ in cols:
                raw = r.values.get(k)
                v = parse_amount(raw, sep) if raw else None
                nums.append(float(v) if v is not None else None)      # số BCTC dưới 2^53: float vẫn đúng từng đơn vị
            ev = [SHORT.get(r.evidence.get(k, {}).get("status"), "") if r.values.get(k) else "" for k, _ in cols]
            ws.append([r.stt, r.code, r.label, r.note, *nums, *ev, r.page])
            row = ws.max_row
            for j, (k, _) in enumerate(cols):
                status = r.evidence.get(k, {}).get("status")
                if status in FILLS and r.values.get(k):
                    ws.cell(row, 5 + j).fill = PatternFill("solid", fgColor=FILLS[status])
                ws.cell(row, 5 + j).number_format = "#,##0;(#,##0)"
        widths = [6, 10, 48, 10, *[18] * len(cols), *[24] * len(cols), 7]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "E3"

    ws = wb.create_sheet("Nguồn", 0)
    lines = [
        ("Mã", info["ticker"]), ("Công ty", info["name"]), ("BCTC", info["title"]), ("Kỳ", info["period_text"]),
        ("Phạm vi", info["scope_text"]), ("Mẫu biểu", d["template"] or ""),
        ("Nguồn", f"{info['source']}: {info['source_url']}"),
        ("Công bố sớm nhất", f"{_when(info['first_at'])} ({info['first_source']})"),
        ("Hệ thống xử lý xong", _when(info["processed_at"])),
        ("Trạng thái", info["status_text"]),
        ("Ràng buộc kế toán bắt buộc", f"{info['gate'][0]}/{info['gate'][1]} đạt"),
        ("File đã đọc", f"{d['chosen']} ({d['page_count']} trang); cột Trang trong các trang tính là số trang của file này"),
        ("", ""),
        ("Cách đọc", "Số chép nguyên văn như in trên BCTC, theo đơn vị ghi ở dòng đầu mỗi trang tính; số trong ngoặc là số "
                     "âm. Cột \"bằng chứng\" cho biết ô được kiểm bằng cách nào (màu giống trang tài liệu trên web)."),
        ("Lưu ý", "Thông tin tổng hợp từ BCTC công bố, không phải khuyến nghị đầu tư. Số do mô hình AI đọc từ ảnh và được "
                  "kiểm bằng ràng buộc kế toán; ô chưa có phép kiểm vẫn có thể sai."),
        ("", ""),
        ("Bằng chứng", "Ý nghĩa"),
    ]
    lines += [(SHORT[k], STATUS_LABELS[k]) for k in STATUS_LABELS if k in SHORT]
    fixes = record.get("manual_fixes", [])
    if fixes:
        lines += [("", ""), ("Ô đã sửa tay sau khi đối chiếu ảnh", f"{len(fixes)} ô")]
        lines += [(f"{m['statement']} dòng {m['row']} {m['column_kind']}", f"\"{m.get('label')}\": {m['before']} → "
                   f"{m['after']} ({m.get('note') or ''})") for m in fixes]
    for a, b in lines:
        ws.append([a, b])
        if a in ("Bằng chứng", "Cách đọc", "Lưu ý") or a.startswith("Ô đã sửa"):
            ws.cell(ws.max_row, 1).font = BOLD
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].alignment = Alignment(vertical="top")
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
